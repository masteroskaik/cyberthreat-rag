import io
import json
from typing import List

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from docx import Document
import markdown
from htmldocx import HtmlToDocx

from api.schemas import ReportItem
from api.auth import get_current_user
from db.connection import get_connection

router = APIRouter(prefix="/reports", tags=["reports"])


def _row_to_report(row) -> ReportItem:
    return ReportItem(
        id=row[0],
        question=row[1],
        answer=row[2],
        sources=row[3] or [],
        created_at=row[4].isoformat(),
    )


@router.get("", response_model=List[ReportItem])
def list_reports(limit: int = 50, current_user: str = Depends(get_current_user)):
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT id, question, answer, sources, created_at "
                "FROM reports ORDER BY created_at DESC LIMIT %s",
                (limit,),
            )
            rows = cur.fetchall()

    return [_row_to_report(r) for r in rows]


@router.get("/{report_id}", response_model=ReportItem)
def get_report(report_id: int, current_user: str = Depends(get_current_user)):
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT id, question, answer, sources, created_at FROM reports WHERE id = %s",
                (report_id,),
            )
            row = cur.fetchone()

    if row is None:
        raise HTTPException(status_code=404, detail="Rapport introuvable")

    return _row_to_report(row)


@router.get("/{report_id}/export/docx")
def export_report_docx(report_id: int, current_user: str = Depends(get_current_user)):
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT id, question, answer, sources, created_at FROM reports WHERE id = %s",
                (report_id,),
            )
            row = cur.fetchone()

    if row is None:
        raise HTTPException(status_code=404, detail="Rapport introuvable")

    report = _row_to_report(row)

    doc = Document()
    doc.add_heading("Rapport CTI — CyberThreat RAG", level=1)
    doc.add_paragraph(f"Généré le : {report.created_at}")
    p = doc.add_paragraph()
    p.add_run(f"Question : {report.question}").bold = True

    doc.add_heading("Analyse", level=2)
    answer_html = markdown.markdown(report.answer, extensions=["tables", "fenced_code"])
    HtmlToDocx().add_html_to_document(answer_html, doc)

    if report.sources:
        doc.add_heading("Sources", level=2)
        for s in report.sources:
            doc.add_paragraph(
                f"{s.get('type', '')} {s.get('id', '')} (score: {s.get('score', 0):.2f})",
                style="List Bullet",
            )

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename=rapport_cti_{report_id}.docx"},
    )
